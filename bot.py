from flask import Flask, request, jsonify
import os
import io
import re
import json
import time
import hashlib
import threading
import requests
import unicodedata

from datetime import datetime, timedelta, timezone, date as dt_date, time as dt_time
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

# =========================
# CONFIGURACIÓN GENERAL
# =========================
WEBEX_TOKEN = os.environ.get("WEBEX_TOKEN", "").strip()            # Token del BOT
WEBEX_USER_TOKEN = os.environ.get("WEBEX_USER_TOKEN", "").strip()  # Token personal para reuniones
WEBEX_API_MESSAGES = "https://webexapis.com/v1/messages"
WEBEX_API_MEETINGS = "https://webexapis.com/v1/meetings"

# Fuente del Excel:
# 1) Si usas SharePoint: EXCEL_SOURCE_URL
# 2) Si usas archivo local: EXCEL_LOCAL_PATH
EXCEL_SOURCE_URL = os.environ.get("EXCEL_SOURCE_URL", "").strip()
EXCEL_LOCAL_PATH = os.environ.get("EXCEL_LOCAL_PATH", "avisos.xlsx").strip()

# Estado interno de avisos enviados
STATE_FILE = os.environ.get("STATE_FILE", "sent_notifications.json").strip()

# Cada cuántos segundos revisará el Excel
CHECK_INTERVAL_SECONDS = int(os.environ.get("CHECK_INTERVAL_SECONDS", "60"))

# Para evitar que mande avisos muy antiguos al arrancar:
# solo enviará filas cuyo horario ya pasó pero no por más de X minutos
ALLOWED_DELAY_MINUTES = int(os.environ.get("ALLOWED_DELAY_MINUTES", "15"))

# Para activar/desactivar el scheduler
RUN_EXCEL_SCHEDULER = os.environ.get("RUN_EXCEL_SCHEDULER", "true").strip().lower() == "true"

TZ = ZoneInfo("America/Lima")
app = Flask(__name__)

# Locks / flags del scheduler
_scheduler_lock = threading.Lock()
_scheduler_started = False
_state_lock = threading.Lock()


# =========================
# HELPERS WEBEX
# =========================
def get_bot_headers():
    """Headers con el token del BOT para enviar mensajes al chat."""
    if not WEBEX_TOKEN:
        raise RuntimeError("Falta WEBEX_TOKEN (Bot) en variables de entorno.")
    return {
        "Authorization": f"Bearer {WEBEX_TOKEN}",
        "Content-Type": "application/json"
    }


# Alias de compatibilidad si en algún lado lo llamabas así
get_webex_headers = get_bot_headers


def get_user_headers():
    """Headers con tu token personal para agendar reuniones."""
    token = WEBEX_USER_TOKEN if WEBEX_USER_TOKEN else WEBEX_TOKEN
    if not token:
        raise RuntimeError("Falta WEBEX_USER_TOKEN o WEBEX_TOKEN.")
    return {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json"
    }


def send_message(room_id: str, text: str):
    """Envía mensaje al room usando el token del bot."""
    data = {"roomId": room_id, "markdown": text}
    try:
        response = requests.post(
            WEBEX_API_MESSAGES,
            headers=get_bot_headers(),
            json=data,
            timeout=15
        )
        if response.status_code in (200, 201):
            return True, response.json()
        return False, response.text
    except Exception as e:
        print(f"❌ Error al enviar mensaje: {e}")
        return False, str(e)


# =========================
# WEBEX MEETINGS
# =========================
def crear_webinar_api(titulo, fecha_inicio_dt, duracion_minutos, panelistas_emails):
    """
    Crea una sesión/reunión en Webex Meetings.
    """
    fecha_utc = fecha_inicio_dt.astimezone(timezone.utc)
    fecha_fin_utc = fecha_utc + timedelta(minutes=duracion_minutos)

    start_iso = fecha_utc.strftime("%Y-%m-%dT%H:%M:%SZ")
    end_iso = fecha_fin_utc.strftime("%Y-%m-%dT%H:%M:%SZ")

    payload = {
        "title": titulo,
        "start": start_iso,
        "end": end_iso
    }

    if panelistas_emails:
        payload["invitees"] = [{"email": email} for email in panelistas_emails if email]

    try:
        response = requests.post(
            WEBEX_API_MEETINGS,
            headers=get_user_headers(),
            json=payload,
            timeout=20
        )

        if response.status_code in (200, 201):
            return True, response.json()

        try:
            err_json = response.json()
        except Exception:
            err_json = {"message": response.text}

        print(f"❌ Respuesta de Webex ({response.status_code}): {err_json}")

        mensaje_error = err_json.get("message", f"Error HTTP {response.status_code}")
        if "errors" in err_json and isinstance(err_json["errors"], list) and err_json["errors"]:
            mensaje_error = err_json["errors"][0].get("description", mensaje_error)

        return False, mensaje_error

    except Exception as e:
        return False, str(e)


def generar_invitacion_circuit(titulo, fecha_str, web_link):
    """Genera invitación con Circuit AI y devuelve (texto, error)."""
    APP_KEY = os.environ.get(
        "CIRCUIT_APP_KEY",
        "egai-prd-operations-108332990-newcontent-1785865382247"
    ).strip()
    CIRCUIT_URL = os.environ.get(
        "CIRCUIT_API_URL",
        "https://circuit.cisco.com/api/v1/chat/completions"
    ).strip()

    headers = {
        "Content-Type": "application/json",
        "appkey": APP_KEY,
        "api-key": APP_KEY
    }

    prompt = (
        f"Redacta una invitación formal para el webinar '{titulo}' "
        f"a realizarse el {fecha_str}. Enlace: {web_link}"
    )

    payload = {
        "agent_id": "bd2530eb-3aeb-4ffc-a6ac-fb748eeb1638",
        "messages": [{"role": "user", "content": prompt}]
    }

    try:
        response = requests.post(CIRCUIT_URL, json=payload, headers=headers, timeout=15)
        if response.status_code == 200:
            res_json = response.json()
            return res_json["choices"][0]["message"]["content"], None
        return None, f"HTTP {response.status_code}: {response.text[:300]}"
    except Exception as e:
        return None, str(e)


# =========================
# MENÚ / COMANDOS
# =========================
def mostrar_menu_principal(room_id):
    menu_texto = (
        "🤖 **¡Hola! Bienvenido al Bot de Gestión de Webex**\n\n"
        "Selecciona una de las siguientes opciones:\n\n"
        "1️⃣ **Crear Webinar / Sesión Webex**\n"
        "2️⃣ **Próximamente: Consultas**\n"
        "3️⃣ **Próximamente: Reportes**\n\n"
        "💡 *Escribe **1** o el comando para crear la sesión.*"
    )
    send_message(room_id, menu_texto)


def mostrar_instrucciones_opcion_1(room_id):
    instrucciones = (
        "📌 **OPCIÓN 1: Crear Webinar / Sesión**\n\n"
        "Envía la información con este formato exacto:\n\n"
        "```text\n"
        "crear webinar;\n"
        "Título: Lanzamiento de Producto 2026;\n"
        "Fecha: 2026-08-15 15:00;\n"
        "Duración: 60;\n"
        "Panelistas: juan@empresa.com, maria@empresa.com\n"
        "```"
    )
    send_message(room_id, instrucciones)


def extraer_datos_webinar(texto_mensaje: str):
    """Parsea título, fecha, duración y panelistas desde el mensaje."""
    lineas = re.split(r'[\n;]+', texto_mensaje)

    datos = {}
    for linea in lineas:
        if ":" in linea:
            clave, valor = linea.split(":", 1)
            clave_clean = (
                clave.strip()
                .lower()
                .replace("í", "i")
                .replace("ó", "o")
            )
            datos[clave_clean] = valor.strip()

    titulo = datos.get("titulo")
    fecha_raw = datos.get("fecha")
    duracion_raw = datos.get("duracion", "30")
    panelistas_raw = datos.get("panelistas", "")

    if not titulo:
        return False, "❌ Falta el campo **Título:** en el mensaje."
    if not fecha_raw:
        return False, "❌ Falta el campo **Fecha:** en el mensaje."

    try:
        fecha_dt = datetime.strptime(fecha_raw, "%Y-%m-%d %H:%M")
        fecha_dt = fecha_dt.replace(tzinfo=TZ)
    except ValueError:
        return False, (
            f"❌ Formato de fecha inválido (`{fecha_raw}`). "
            f"Usa: `YYYY-MM-DD HH:MM` (Ejemplo: `2026-08-04 15:00`)."
        )

    try:
        duracion = int(re.sub(r'\D', '', duracion_raw))
    except ValueError:
        duracion = 30

    panelistas = [e.strip() for e in re.split(r'[,;\s]+', panelistas_raw) if "@" in e]

    return True, {
        "titulo": titulo,
        "fecha_dt": fecha_dt,
        "duracion": duracion,
        "panelistas": panelistas
    }


# =========================
# EXCEL / AVISOS
# =========================
def normalizar_texto(texto: str) -> str:
    if texto is None:
        return ""
    texto = str(texto).strip()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.lower()
    texto = re.sub(r"\s+", "", texto)
    return texto


def build_sharepoint_download_url(url: str) -> str:
    """Intenta convertir el link compartido en descarga directa."""
    if not url:
        return url
    if "download=1" in url.lower():
        return url
    separador = "&" if "?" in url else "?"
    return f"{url}{separador}download=1"


def obtener_fuente_excel():
    """
    Devuelve:
    - un BytesIO si viene por URL
    - una ruta local si viene por archivo
    """
    if EXCEL_SOURCE_URL:
        excel_url = build_sharepoint_download_url(EXCEL_SOURCE_URL)
        response = requests.get(excel_url, timeout=30)

        if response.status_code != 200:
            raise RuntimeError(
                f"No se pudo descargar el Excel desde URL. "
                f"HTTP {response.status_code}: {response.text[:200]}"
            )

        content_type = response.headers.get("Content-Type", "")
        if "html" in content_type.lower():
            raise RuntimeError(
                "La URL devolvió HTML en vez de un archivo Excel. "
                "Probablemente el link de SharePoint requiere autenticación "
                "o no es una descarga directa."
            )

        return io.BytesIO(response.content)

    if EXCEL_LOCAL_PATH and os.path.exists(EXCEL_LOCAL_PATH):
        return EXCEL_LOCAL_PATH

    raise RuntimeError(
        "No hay fuente de Excel disponible. Configura EXCEL_SOURCE_URL "
        "o sube un archivo y define EXCEL_LOCAL_PATH."
    )


def parsear_fecha_excel(valor):
    if valor is None or str(valor).strip() == "":
        raise ValueError("Fecha vacía")

    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, dt_date):
        return valor

    if isinstance(valor, (int, float)):
        try:
            dt = from_excel(valor)
            if isinstance(dt, datetime):
                return dt.date()
            if isinstance(dt, dt_date):
                return dt
        except Exception:
            pass

    texto = str(valor).strip()
    formatos = [
        "%m/%d/%Y",   # 11/21/2025
        "%m/%d/%y",
        "%d/%m/%Y",
        "%Y-%m-%d",
    ]
    for fmt in formatos:
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue

    raise ValueError(f"Fecha no reconocida: {valor}")


def parsear_hora_excel(valor):
    if valor is None or str(valor).strip() == "":
        raise ValueError("Hora vacía")

    if isinstance(valor, datetime):
        return valor.time().replace(tzinfo=None)

    if isinstance(valor, dt_time):
        return valor.replace(tzinfo=None)

    if isinstance(valor, (int, float)):
        try:
            dt = from_excel(valor)
            if isinstance(dt, datetime):
                return dt.time().replace(tzinfo=None)
        except Exception:
            pass

        total_seconds = int(round(float(valor) * 24 * 60 * 60))
        total_seconds = total_seconds % (24 * 60 * 60)
        horas = total_seconds // 3600
        minutos = (total_seconds % 3600) // 60
        segundos = total_seconds % 60
        return dt_time(hour=horas, minute=minutos, second=segundos)

    texto = str(valor).strip()
    formatos = [
        "%I:%M:%S %p",  # 4:16:00 PM
        "%I:%M %p",
        "%H:%M:%S",
        "%H:%M"
    ]
    for fmt in formatos:
        try:
            return datetime.strptime(texto, fmt).time()
        except ValueError:
            continue

    raise ValueError(f"Hora no reconocida: {valor}")


def cargar_estado_enviados():
    with _state_lock:
        if not os.path.exists(STATE_FILE):
            return {}

        try:
            with open(STATE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}


def guardar_estado_enviados(estado):
    with _state_lock:
        tmp_path = f"{STATE_FILE}.tmp"
        with open(tmp_path, "w", encoding="utf-8") as f:
            json.dump(estado, f, ensure_ascii=False, indent=2)
        os.replace(tmp_path, STATE_FILE)


def generar_clave_aviso(aviso):
    base = (
        f"{aviso['fila']}|"
        f"{aviso['fecha_envio'].isoformat()}|"
        f"{aviso['room_id']}|"
        f"{aviso['mensaje']}"
    )
    return hashlib.sha256(base.encode("utf-8")).hexdigest()


def leer_avisos_excel():
    """
    Lee el Excel con columnas:
    Fecha | Hora | RoomID | Mensaje
    """
    fuente = obtener_fuente_excel()
    wb = load_workbook(fuente, data_only=True)
    ws = wb.active

    encabezados_raw = [c.value for c in ws[1]]
    encabezados = {}

    for idx, valor in enumerate(encabezados_raw):
        clave = normalizar_texto(valor)
        if clave:
            encabezados[clave] = idx

    requeridos = ["fecha", "hora", "roomid", "mensaje"]
    faltantes = [x for x in requeridos if x not in encabezados]
    if faltantes:
        raise RuntimeError(
            f"Faltan columnas requeridas en Excel: {faltantes}. "
            f"Deben existir: Fecha, Hora, RoomID, Mensaje"
        )

    avisos = []
    errores = []

    idx_fecha = encabezados["fecha"]
    idx_hora = encabezados["hora"]
    idx_roomid = encabezados["roomid"]
    idx_mensaje = encabezados["mensaje"]

    for fila_num, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            fecha_val = fila[idx_fecha] if idx_fecha < len(fila) else None
            hora_val = fila[idx_hora] if idx_hora < len(fila) else None
            room_id_val = fila[idx_roomid] if idx_roomid < len(fila) else None
            mensaje_val = fila[idx_mensaje] if idx_mensaje < len(fila) else None

            if not fecha_val and not hora_val and not room_id_val and not mensaje_val:
                continue

            room_id = str(room_id_val).strip() if room_id_val is not None else ""
            mensaje = str(mensaje_val).strip() if mensaje_val is not None else ""

            if not room_id or not mensaje:
                errores.append({
                    "fila": fila_num,
                    "error": "RoomID o Mensaje vacío"
                })
                continue

            fecha = parsear_fecha_excel(fecha_val)
            hora = parsear_hora_excel(hora_val)

            fecha_envio = datetime.combine(fecha, hora).replace(tzinfo=TZ)

            avisos.append({
                "fila": fila_num,
                "fecha_envio": fecha_envio,
                "room_id": room_id,
                "mensaje": mensaje
            })

        except Exception as e:
            errores.append({
                "fila": fila_num,
                "error": str(e)
            })

    return avisos, errores


def procesar_avisos_pendientes():
    """
    Revisa el Excel y envía avisos pendientes.
    """
    ahora = datetime.now(TZ)
    limite_atras = ahora - timedelta(minutes=ALLOWED_DELAY_MINUTES)

    avisos, errores = leer_avisos_excel()
    estado = cargar_estado_enviados()

    enviados = []
    omitidos = []
    expirados = []

    for aviso in avisos:
        clave = generar_clave_aviso(aviso)

        if clave in estado:
            omitidos.append({
                "fila": aviso["fila"],
                "motivo": "ya_enviado"
            })
            continue

        if aviso["fecha_envio"] > ahora:
            omitidos.append({
                "fila": aviso["fila"],
                "motivo": "todavia_no_toca"
            })
            continue

        if aviso["fecha_envio"] < limite_atras:
            expirados.append({
                "fila": aviso["fila"],
                "motivo": "demasiado_antiguo",
                "programado_para": aviso["fecha_envio"].isoformat()
            })
            continue

        ok, detalle = send_message(aviso["room_id"], aviso["mensaje"])
        if ok:
            estado[clave] = {
                "fila": aviso["fila"],
                "roomId": aviso["room_id"],
                "mensaje": aviso["mensaje"],
                "programado_para": aviso["fecha_envio"].isoformat(),
                "enviado_en": ahora.isoformat()
            }
            enviados.append({
                "fila": aviso["fila"],
                "roomId": aviso["room_id"],
                "mensaje": aviso["mensaje"]
            })
        else:
            errores.append({
                "fila": aviso["fila"],
                "error": f"Fallo al enviar a Webex: {detalle}"
            })

    guardar_estado_enviados(estado)

    return {
        "ok": True,
        "ahora": ahora.isoformat(),
        "total_avisos_leidos": len(avisos),
        "enviados": enviados,
        "omitidos": omitidos,
        "expirados": expirados,
        "errores": errores
    }


def scheduler_avisos_loop():
    print("⏰ Scheduler de avisos iniciado.")
    while True:
        try:
            resumen = procesar_avisos_pendientes()
            enviados = len(resumen.get("enviados", []))
            if enviados > 0:
                print(f"✅ Avisos enviados en este ciclo: {enviados}")
        except Exception as e:
            print(f"❌ Error en scheduler de avisos: {e}")

        time.sleep(CHECK_INTERVAL_SECONDS)


def iniciar_scheduler_si_corresponde():
    global _scheduler_started

    if not RUN_EXCEL_SCHEDULER:
        print("ℹ️ Scheduler de Excel desactivado por RUN_EXCEL_SCHEDULER=false")
        return

    with _scheduler_lock:
        if _scheduler_started:
            return

        hilo = threading.Thread(target=scheduler_avisos_loop, daemon=True)
        hilo.start()
        _scheduler_started = True
        print("✅ Scheduler de Excel lanzado.")


# =========================
# DEBUG EXCEL
# =========================
@app.route("/debug/excel", methods=["GET"])
def debug_excel():
    try:
        avisos, errores = leer_avisos_excel()
        preview = []
        for aviso in avisos[:10]:
            preview.append({
                "fila": aviso["fila"],
                "fecha_envio": aviso["fecha_envio"].isoformat(),
                "room_id_preview": aviso["room_id"][:35] + ("..." if len(aviso["room_id"]) > 35 else ""),
                "mensaje": aviso["mensaje"]
            })

        return jsonify({
            "ok": True,
            "timezone": str(TZ),
            "check_interval_seconds": CHECK_INTERVAL_SECONDS,
            "allowed_delay_minutes": ALLOWED_DELAY_MINUTES,
            "excel_source_url_configured": bool(EXCEL_SOURCE_URL),
            "excel_local_path": EXCEL_LOCAL_PATH,
            "total_avisos": len(avisos),
            "preview": preview,
            "errores": errores[:20]
        }), 200

    except Exception as e:
        return jsonify({
            "ok": False,
            "error": str(e)
        }), 500


@app.route("/debug/run-notifications", methods=["GET", "POST"])
def debug_run_notifications():
    try:
        resumen = procesar_avisos_pendientes()
        return jsonify(resumen), 200
    except Exception as e:
        return jsonify({
            "ok": False,
            "error": str(e)
        }), 500


# =========================
# WEBHOOK DE WEBEX
# =========================
@app.route("/webhook", methods=["POST"])
def webhook():
    data = request.json or {}
    if "data" not in data:
        return "ok", 200

    msg_id = data["data"].get("id")
    if not msg_id:
        return "ok", 200

    try:
        res = requests.get(
            f"{WEBEX_API_MESSAGES}/{msg_id}",
            headers=get_bot_headers(),
            timeout=15
        )
        if res.status_code != 200:
            return "ok", 200

        msg = res.json()
        raw_text = (msg.get("text") or msg.get("markdown") or "").strip()
        texto_lower = raw_text.lower()
        room = msg.get("roomId")
        sender = msg.get("personEmail", "")

        # Ignorar mensajes del propio bot u otros bots
        if not sender or not raw_text or sender.endswith("@webex.bot"):
            return "ok", 200

        print(f"📩 [{sender}]: {raw_text}")

        if texto_lower in ["hola", "menu", "menú", "opciones", "start", "ayuda"]:
            mostrar_menu_principal(room)

        elif texto_lower == "1":
            mostrar_instrucciones_opcion_1(room)

        elif texto_lower in ["debug excel", "ver excel"]:
            avisos, errores = leer_avisos_excel()
            send_message(
                room,
                f"📊 **Excel leído correctamente**\n\n"
                f"- Avisos válidos: **{len(avisos)}**\n"
                f"- Errores detectados: **{len(errores)}**"
            )

        elif texto_lower in ["ejecutar avisos", "probar avisos"]:
            resumen = procesar_avisos_pendientes()
            send_message(
                room,
                f"⏱️ **Ejecución manual completada**\n\n"
                f"- Leídos: **{resumen['total_avisos_leidos']}**\n"
                f"- Enviados: **{len(resumen['enviados'])}**\n"
                f"- Errores: **{len(resumen['errores'])}**\n"
                f"- Expirados: **{len(resumen['expirados'])}**"
            )

        elif "crear webinar" in texto_lower:
            exito, resultado = extraer_datos_webinar(raw_text)

            if not exito:
                send_message(room, resultado)
            else:
                datos = resultado
                send_message(room, f"⌛ Creando sesión **'{datos['titulo']}'**...")

                creado, res_api = crear_webinar_api(
                    titulo=datos["titulo"],
                    fecha_inicio_dt=datos["fecha_dt"],
                    duracion_minutos=datos["duracion"],
                    panelistas_emails=datos["panelistas"]
                )

                if creado:
                    web_link = res_api.get("webLink", "No disponible")
                    fecha_formateada = datos["fecha_dt"].strftime("%Y-%m-%d %H:%M")

                    send_message(
                        room,
                        f"✅ **¡Sesión creada con éxito!**\n\n"
                        f"📌 **Título:** {datos['titulo']}\n"
                        f"🔗 **Enlace:** {web_link}\n\n"
                        f"🤖 *Generando borrador de invitación con Circuit AI...*"
                    )

                    borrador, error_circuit = generar_invitacion_circuit(
                        titulo=datos["titulo"],
                        fecha_str=fecha_formateada,
                        web_link=web_link
                    )

                    if borrador:
                        send_message(room, f"📧 **Borrador de Correo Generado por Circuit:**\n\n{borrador}")
                    else:
                        send_message(room, f"⚠️ **Detalle del fallo de Circuit AI:**\n`{error_circuit}`")

                else:
                    send_message(room, f"❌ **Error al crear sesión en Webex:** {res_api}")

        else:
            send_message(
                room,
                "❓ Comando no reconocido.\n\n"
                "Escribe **menu** para ver las opciones.\n"
                "O usa **debug excel** para validar el archivo."
            )

    except Exception as e:
        print("❌ Error general en webhook:", e)

    return "ok", 200


# =========================
# HEALTHCHECK
# =========================
@app.route("/ping", methods=["GET", "HEAD"])
def ping():
    return "Servidor activo", 200


# Iniciar scheduler al cargar la app
iniciar_scheduler_si_corresponde()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")))































